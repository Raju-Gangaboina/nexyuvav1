"""
Live E2E smoke test for Module 5 (Recruiter Filters). Registers 2 students
(one matching a degree/location filter, one not), one org, posts an
internship, then exercises match_threshold / degree / location / skills /
only_untapped against the real running server + real DB.
"""
import asyncio
import uuid

import httpx
from docx import Document

BASE = "http://localhost:8000/api/v1"


def make_resume(path: str, skills_line: str):
    doc = Document()
    doc.add_paragraph("Test Candidate")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph(skills_line)
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("B.Tech Computer Science, CGPA 8.2, 2027")
    doc.save(path)


async def register_and_upload(client, tag, email_prefix, skills_line, degree, location):
    r = await client.post(f"{BASE}/auth/register/student", json={
        "email": f"{email_prefix}_{tag}@test.com",
        "password": "TestPass123!",
        "full_name": f"{email_prefix.title()} Candidate",
    })
    r.raise_for_status()
    token = r.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    resume_path = f"test_filters_resume_{email_prefix}_{tag}.docx"
    make_resume(resume_path, skills_line)
    with open(resume_path, "rb") as f:
        r = await client.post(
            f"{BASE}/students/me/resume", headers=headers,
            files={"file": (resume_path, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    r.raise_for_status()

    r = await client.patch(f"{BASE}/students/me", headers=headers, json={
        "degree": degree, "location": location, "year": 2027,
    })
    r.raise_for_status()
    return headers


async def main():
    tag = uuid.uuid4().hex[:8]

    async with httpx.AsyncClient(timeout=60) as client:
        strong_headers = await register_and_upload(
            client, tag, "strong", "Python, Django, PostgreSQL",
            degree="B.Tech Computer Science", location="Bangalore",
        )
        weak_headers = await register_and_upload(
            client, tag, "weak", "Excel, Communication",
            degree="B.Com", location="Delhi",
        )

        r = await client.post(f"{BASE}/auth/register/organization", json={
            "email": f"org_filt_{tag}@test.com",
            "password": "TestPass123!",
            "organization_name": f"Filter Test Org {tag}",
        })
        r.raise_for_status()
        org_headers = {"Authorization": f"Bearer {r.json()['access_token']}"}

    from sqlalchemy import select
    from app.core.database import AsyncSessionLocal
    from app.core.models import Organization, OrgVerificationStatus, StudentProfile, User

    async with AsyncSessionLocal() as db:
        result = await db.execute(select(User).where(User.email == f"org_filt_{tag}@test.com"))
        org_user = result.scalar_one()
        result = await db.execute(select(Organization).where(Organization.user_id == org_user.id))
        org = result.scalar_one()
        org.is_verified = True
        org.verification_status = OrgVerificationStatus.APPROVED
        await db.commit()

        result = await db.execute(select(User).where(User.email == f"strong_{tag}@test.com"))
        strong_user = result.scalar_one()
        result = await db.execute(select(StudentProfile).where(StudentProfile.user_id == strong_user.id))
        strong_profile = result.scalar_one()
        strong_id = str(strong_profile.id)

        result = await db.execute(select(User).where(User.email == f"weak_{tag}@test.com"))
        weak_user = result.scalar_one()
        result = await db.execute(select(StudentProfile).where(StudentProfile.user_id == weak_user.id))
        weak_profile = result.scalar_one()
        weak_id = str(weak_profile.id)

    async with httpx.AsyncClient(timeout=60) as client:
        r = await client.post(f"{BASE}/organizations/internships", headers=org_headers, json={
            "title": "Backend Intern",
            "domain": "Software Development",
            "required_skills": ["python", "django", "postgresql"],
            "description": "Backend internship building APIs in Python/Django.",
            "duration_weeks": 12,
            "location": "Remote",
            "work_mode": "remote",
        })
        r.raise_for_status()
        internship_id = r.json()["id"]

        async def sourcing(params):
            r = await client.get(
                f"{BASE}/organizations/internships/{internship_id}/sourcing",
                headers=org_headers,
                params={"top_k": 10, "with_ai_pitch": False, **params},
            )
            r.raise_for_status()
            return {c["student_id"] for c in r.json()}

        baseline = await sourcing({})
        print("BASELINE_HAS_STRONG:", strong_id in baseline)
        print("BASELINE_HAS_WEAK:", weak_id in baseline)

        by_degree = await sourcing({"degree": "Computer Science"})
        print("DEGREE_FILTER -> strong in:", strong_id in by_degree, "| weak in:", weak_id in by_degree)
        assert strong_id in by_degree and weak_id not in by_degree, "degree filter failed"

        by_location = await sourcing({"location": "bangalore"})
        print("LOCATION_FILTER (case-insensitive) -> strong in:", strong_id in by_location, "| weak in:", weak_id in by_location)
        assert strong_id in by_location and weak_id not in by_location, "location filter failed"

        by_skills = await sourcing({"skills": "python,django"})
        print("SKILLS_FILTER (all-of) -> strong in:", strong_id in by_skills, "| weak in:", weak_id in by_skills)
        assert strong_id in by_skills and weak_id not in by_skills, "skills filter failed"

        by_threshold_high = await sourcing({"match_threshold": 90})
        print("MATCH_THRESHOLD=90 -> strong in:", strong_id in by_threshold_high, "| weak in:", weak_id in by_threshold_high)
        assert weak_id not in by_threshold_high, "match_threshold should exclude the weak/no-overlap candidate"

        # Apply with the strong candidate, then confirm only_untapped drops them
        r = await client.post(f"{BASE}/students/me/applications", headers=strong_headers, json={
            "internship_id": internship_id,
        })
        r.raise_for_status()
        by_untapped = await sourcing({"only_untapped": True})
        print("ONLY_UNTAPPED (after strong applied) -> strong in:", strong_id in by_untapped)
        assert strong_id not in by_untapped, "only_untapped should exclude a candidate who just applied"

    print("ALL_FILTER_ASSERTIONS_PASSED")


asyncio.run(main())
