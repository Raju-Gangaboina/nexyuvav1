"""
Live E2E smoke test for the Certifications evidence signal + Internship
History context bonus (Module 1 gap-closure follow-up).

Registers a student with a resume containing an AWS certification and a
Python project, verifies an org, posts two internships in the SAME domain,
has the student apply to the first, then sources the second and confirms:
  - "aws" skill_confidence includes "Backed by a listed certification"
  - the explanation includes the internship-history context bonus phrase

Run against the already-running `uvicorn --reload` dev server on :8000.
"""
import asyncio
import uuid

import httpx
from docx import Document

BASE = "http://localhost:8000/api/v1"


def make_resume(path: str):
    doc = Document()
    doc.add_paragraph("Cert History Test Candidate")
    doc.add_paragraph("test.candidate@example.com")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Python, AWS, Docker")
    doc.add_paragraph("PROJECTS")
    doc.add_paragraph("Cloud Deploy Tool - Automated deployment tool (Python, AWS, Docker)")
    doc.add_paragraph("Built a Python-based deployment automation tool using AWS Lambda.")
    doc.add_paragraph("CERTIFICATIONS")
    doc.add_paragraph("AWS Certified Cloud Practitioner - Amazon Web Services, 2025")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("B.Tech Computer Science, CGPA 8.4, 2026")
    doc.save(path)


async def main():
    tag = uuid.uuid4().hex[:8]
    resume_path = f"test_cert_resume_{tag}.docx"
    make_resume(resume_path)

    async with httpx.AsyncClient(timeout=60) as client:
        r = await client.post(f"{BASE}/auth/register/student", json={
            "email": f"certstudent_{tag}@test.com",
            "password": "TestPass123!",
            "full_name": "Cert History Test Candidate",
        })
        r.raise_for_status()
        student_token = r.json()["access_token"]
        student_headers = {"Authorization": f"Bearer {student_token}"}

        with open(resume_path, "rb") as f:
            r = await client.post(
                f"{BASE}/students/me/resume",
                headers=student_headers,
                files={"file": (resume_path, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        r.raise_for_status()
        parsed = r.json()
        print("PARSED_CERTIFICATIONS:", parsed.get("certifications"))
        print("PARSED_PROJECTS:", [p["title"] for p in parsed.get("projects", [])])

        r = await client.get(f"{BASE}/students/me", headers=student_headers)
        r.raise_for_status()
        profile = r.json()
        print("PERSISTED_CERTIFICATIONS_ON_PROFILE:", profile.get("certifications"))

        r = await client.post(f"{BASE}/auth/register/organization", json={
            "email": f"certorg_{tag}@test.com",
            "password": "TestPass123!",
            "organization_name": f"Cert Org {tag}",
        })
        r.raise_for_status()
        org_token = r.json()["access_token"]
        org_headers = {"Authorization": f"Bearer {org_token}"}

    from sqlalchemy import select
    from app.core.database import AsyncSessionLocal
    from app.core.models import Organization, OrgVerificationStatus, StudentProfile, User

    async with AsyncSessionLocal() as db:
        result = await db.execute(select(User).where(User.email == f"certorg_{tag}@test.com"))
        org_user = result.scalar_one()
        result = await db.execute(select(Organization).where(Organization.user_id == org_user.id))
        org = result.scalar_one()
        org.is_verified = True
        org.verification_status = OrgVerificationStatus.APPROVED
        await db.commit()

        result = await db.execute(select(User).where(User.email == f"certstudent_{tag}@test.com"))
        student_user = result.scalar_one()
        result = await db.execute(select(StudentProfile).where(StudentProfile.user_id == student_user.id))
        student_profile = result.scalar_one()
        student_row_id = str(student_profile.id)
        print("DB_CERTIFICATIONS_COLUMN:", student_profile.certifications)

    async with httpx.AsyncClient(timeout=60) as client:
        # Two internships in the SAME domain -- apply to the first, source the second.
        r = await client.post(f"{BASE}/organizations/internships", headers=org_headers, json={
            "title": "Cloud Intern (First)",
            "domain": "DevOps",
            "required_skills": ["aws", "python"],
            "description": "Cloud automation internship.",
            "duration_weeks": 12,
            "location": "Remote",
            "work_mode": "remote",
        })
        r.raise_for_status()
        internship_1_id = r.json()["id"]

        r = await client.post(f"{BASE}/organizations/internships", headers=org_headers, json={
            "title": "Cloud Intern (Second)",
            "domain": "DevOps",
            "required_skills": ["aws", "python", "docker"],
            "description": "Second cloud automation internship, same domain.",
            "duration_weeks": 12,
            "location": "Remote",
            "work_mode": "remote",
        })
        r.raise_for_status()
        internship_2_id = r.json()["id"]

        # Student applies to the FIRST internship only.
        r = await client.post(
            f"{BASE}/students/me/applications",
            headers=student_headers,
            json={"internship_id": internship_1_id},
        )
        print("APPLY_STATUS:", r.status_code, r.text[:200] if r.status_code >= 400 else "OK")

        # Source the SECOND internship -- should show internship-history bonus.
        r = await client.get(
            f"{BASE}/organizations/internships/{internship_2_id}/sourcing",
            headers=org_headers,
            params={"top_k": 10, "with_ai_pitch": False},
        )
        r.raise_for_status()
        candidates = r.json()

    match = next((c for c in candidates if c["student_id"] == student_row_id), None)
    print("CANDIDATE_FOUND:", match is not None)
    if match:
        aws_conf = next((sc for sc in match["skill_confidence"] if sc["skill"] == "aws"), None)
        print("AWS_SKILL_CONFIDENCE:", aws_conf)
        assert aws_conf is not None, "aws should be a matched skill"
        assert "Backed by a listed certification" in aws_conf["evidence"], "certification evidence missing"
        print("CERTIFICATION_EVIDENCE_CONFIRMED")

        print("EXPLANATION:", match["explanation"])
        assert "similar roles before" in match["explanation"], "internship history bonus missing from explanation"
        print("INTERNSHIP_HISTORY_BONUS_CONFIRMED")
        print("ALL_ASSERTIONS_PASSED")
    else:
        print("ALL_CANDIDATE_IDS:", [c["student_id"] for c in candidates])


asyncio.run(main())
